from flask import Flask, render_template, request, jsonify
from werkzeug.utils import secure_filename
import os
import re
import math
from collections import Counter

app = Flask(__name__)

# ============================================================
# CONFIGURATION
# ============================================================

UPLOAD_FOLDER = "uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

ALLOWED_EXTENSIONS = {"pdf", "docx", "doc"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_SIZE

# Project 1 is intentionally RESUME-ONLY.
# Job-description matching belongs to Project 2.
#
# IMPORTANT:
# All weights below add up to exactly 100.
CRITERIA_WEIGHTS = {
    "contact_information": 6,
    "resume_parsing": 6,
    "job_description_matching": 9,  # N/A in Project 1
    "keyword_matching": 6,
    "skills_analysis": 7,
    "resume_sections": 6,
    "experience_matching": 6,
    "project_relevance": 7,
    "education": 5,
    "certifications": 4,
    "achievements_metrics": 5,
    "action_verbs": 5,
    "ats_formatting": 5,
    "readability": 4,
    "keyword_context": 5,
    "missing_requirements": 4,  # N/A without a JD
    "potential_red_flags": 5,
    "content_consistency": 5,
}

assert sum(CRITERIA_WEIGHTS.values()) == 100

CRITERION_LABELS = {
    "contact_information": "Contact Information",
    "resume_parsing": "Resume Parsing",
    "job_description_matching": "Job Description Matching",
    "keyword_matching": "Keyword Matching",
    "skills_analysis": "Skills Analysis",
    "resume_sections": "Resume Sections",
    "experience_matching": "Experience Matching",
    "project_relevance": "Project Relevance",
    "education": "Education",
    "certifications": "Certifications",
    "achievements_metrics": "Achievements & Metrics",
    "action_verbs": "Action Verbs",
    "ats_formatting": "ATS Formatting",
    "readability": "Readability",
    "keyword_context": "Keyword Context",
    "missing_requirements": "Missing Requirements",
    "potential_red_flags": "Potential Red Flags",
    "content_consistency": "Content Consistency",
}

# ============================================================
# SKILLS / PATTERNS
# ============================================================

SKILLS = [
    "python", "java", "c++", "c", "javascript", "typescript",
    "html", "css", "react", "node.js", "flask", "django",
    "sql", "mysql", "postgresql", "mongodb", "git", "github",
    "docker", "aws", "azure", "machine learning", "deep learning",
    "data science", "data analysis", "pandas", "numpy",
    "scikit-learn", "tensorflow", "pytorch", "artificial intelligence",
    "power bi", "tableau", "excel", "statistics", "nlp",
    "computer vision", "rest api", "api", "linux", "devops",
    "kubernetes", "spark", "hadoop", "agile", "jira", "jenkins",
    "communication", "leadership", "problem solving",
    "data visualization", "matplotlib", "seaborn", "fastapi",
    "postman", "redis", "gitlab", "bitbucket", "oracle",
    "sqlite", "scipy", "keras", "pyspark", "airflow"
]

ACTION_WORDS = [
    "developed", "created", "implemented", "designed", "built",
    "managed", "analyzed", "improved", "led", "optimized",
    "automated", "engineered", "deployed", "tested", "maintained",
    "delivered", "integrated", "configured", "debugged", "refactored",
    "streamlined", "evaluated", "researched", "trained", "predicted",
    "visualized", "documented", "collaborated", "presented"
]

SECTION_ALIASES = {
    "summary": [
        "summary", "professional summary", "profile",
        "career objective", "objective"
    ],
    "experience": [
        "experience", "work experience", "professional experience",
        "employment", "internship", "internships"
    ],
    "education": [
        "education", "academic background",
        "academic qualifications"
    ],
    "skills": [
        "skills", "technical skills", "core skills",
        "technical competencies", "technologies"
    ],
    "projects": [
        "projects", "academic projects", "personal projects",
        "project experience"
    ],
    "certifications": [
        "certifications", "certificates", "professional certifications"
    ],
    "achievements": [
        "achievements", "accomplishments", "awards"
    ],
    "languages": [
        "languages", "language skills"
    ]
}

DEGREE_TERMS = [
    "b.tech", "btech", "b.e.", "be ", "bachelor",
    "m.tech", "mtech", "m.s.", "master", "mba", "ph.d", "phd",
    "b.sc", "bsc", "m.sc", "msc", "degree", "university", "college"
]

DATE_PATTERN = re.compile(
    r"\b(?:19|20)\d{2}\b|"
    r"\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)"
    r"(?:uary|ruary|ch|il|e|y|ust|tember|ober|ember)?\s+"
    r"(?:19|20)\d{2}\b",
    re.I
)

EMAIL_PATTERN = re.compile(
    r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b"
)

PHONE_PATTERNS = [
    re.compile(r"(?:\+91[\s.-]?)?[6-9]\d{9}\b"),
    re.compile(r"\b(?:\+\d{1,3}[\s.-]?)?\d{10,12}\b")
]

URL_PATTERN = re.compile(
    r"https?://[^\s]+|www\.[^\s]+",
    re.I
)

# ============================================================
# GENERAL HELPERS
# ============================================================

def allowed_file(filename):
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS
    )


def normalize_text(text):
    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def unique_keep_order(items):
    seen = set()
    result = []

    for item in items:
        key = item.lower() if isinstance(item, str) else item

        if key not in seen:
            seen.add(key)
            result.append(item)

    return result


def clamp(value, low=0.0, high=1.0):
    return max(low, min(high, value))


def make_result(
    key,
    score,
    evidence,
    explanation,
    status=None
):
    maximum = CRITERIA_WEIGHTS[key]

    if status is None:
        ratio = score / maximum if maximum else 0

        if ratio >= 0.80:
            status = "Strong"
        elif ratio >= 0.55:
            status = "Good"
        elif ratio > 0:
            status = "Needs Improvement"
        else:
            status = "Missing"

    return {
        "key": key,
        "criterion": CRITERION_LABELS[key],
        "score": round(score, 2),
        "max_score": maximum,
        "status": status,
        "evidence": (
            evidence
            if evidence
            else ["No strong evidence detected."]
        ),
        "explanation": explanation,
    }


def na_result(key, reason):
    return make_result(
        key,
        0,
        ["Not evaluated in Project 1."],
        reason,
        status="N/A"
    )


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================

def extract_pdf_text(filepath):
    try:
        from pypdf import PdfReader

        reader = PdfReader(filepath)
        pages = []

        for page in reader.pages:
            page_text = page.extract_text() or ""
            pages.append(page_text)

        return "\n".join(pages), len(reader.pages)

    except Exception as exc:
        print("PDF extraction error:", exc)
        return "", 0


def extract_docx_text(filepath):
    try:
        from docx import Document

        document = Document(filepath)

        paragraphs = [
            p.text
            for p in document.paragraphs
            if p.text.strip()
        ]

        # Include table text because resumes sometimes put important
        # content in tables.
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                if any(cells):
                    paragraphs.append(" | ".join(cells))

        return "\n".join(paragraphs), 1

    except Exception as exc:
        print("DOCX extraction error:", exc)
        return "", 0


def extract_resume_text(filepath, filename):
    extension = filename.rsplit(".", 1)[1].lower()

    if extension == "pdf":
        return extract_pdf_text(filepath)

    if extension == "docx":
        return extract_docx_text(filepath)

    # Legacy .doc parsing is intentionally not faked.
    # Convert .doc to DOCX or PDF for reliable extraction.
    if extension == "doc":
        return "", 0

    return "", 0


# ============================================================
# TEXT ANALYSIS
# ============================================================

def find_skills(text):
    text_lower = text.lower()
    found = []

    for skill in SKILLS:
        pattern = (
            r"(?<!\w)"
            + re.escape(skill.lower())
            + r"(?!\w)"
        )

        if re.search(pattern, text_lower):
            found.append(skill)

    return found


def detect_sections(text):
    text_lower = text.lower()
    detected = {}
    evidence = {}

    for section, aliases in SECTION_ALIASES.items():

        detected[section] = False
        evidence[section] = None

        for alias in aliases:
            pattern = (
                r"(?<!\w)"
                + re.escape(alias)
                + r"(?!\w)"
            )

            if re.search(pattern, text_lower):
                detected[section] = True
                evidence[section] = alias
                break

    return detected, evidence


def extract_contact_details(text):
    email_match = EMAIL_PATTERN.search(text)
    phone_match = None

    for pattern in PHONE_PATTERNS:
        phone_match = pattern.search(text)

        if phone_match:
            break

    lower = text.lower()

    return {
        "email": bool(email_match),
        "phone": bool(phone_match),
        "linkedin": "linkedin.com" in lower,
        "github": "github.com" in lower,
        "portfolio": any(
            token in lower
            for token in [
                "portfolio",
                "behance.net",
                "dribbble.com"
            ]
        ),
    }


def count_action_verbs(text):
    lower = text.lower()
    found = []

    for word in ACTION_WORDS:

        if re.search(
            r"(?<!\w)"
            + re.escape(word)
            + r"(?!\w)",
            lower
        ):
            found.append(word)

    return unique_keep_order(found)


def extract_metrics(text):
    patterns = [
        r"\b\d+(?:\.\d+)?\s*%\b",

        r"\b\d+(?:,\d{3})*(?:\+)?\s*"
        r"(?:users|customers|records|employees|"
        r"projects|transactions|rows|requests|datasets)\b",

        r"\b\d+(?:\.\d+)?x\b",

        r"\$\s?\d+(?:,\d{3})*(?:\.\d+)?",

        r"\b(?:increased|decreased|reduced|improved|"
        r"saved|grew|cut|boosted)\b"
        r".{0,80}"
        r"\b\d+(?:\.\d+)?%?"
    ]

    matches = []

    for pattern in patterns:
        matches.extend(
            re.findall(
                pattern,
                text,
                flags=re.I | re.S
            )
        )

    return unique_keep_order(
        [
            m.strip()
            for m in matches
            if m.strip()
        ]
    )


def extract_years(text):
    return unique_keep_order(
        DATE_PATTERN.findall(text)
    )


def split_lines(text):
    return [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]


def bullet_lines(text):
    lines = split_lines(text)

    return [
        line
        for line in lines
        if re.match(
            r"^(?:[-*•▪◦]|\d+[.)])\s+",
            line
        )
    ]


def word_count(text):
    return len(
        re.findall(
            r"\b[\w+#.-]+\b",
            text
        )
    )


def sentence_count(text):
    sentences = re.split(
        r"(?<=[.!?])\s+|\n+",
        text.strip()
    )

    return max(
        1,
        len(
            [
                s
                for s in sentences
                if s.strip()
            ]
        )
    )


# ============================================================
# 1. CONTACT INFORMATION
# ============================================================

def check_contact_information(text):

    details = extract_contact_details(text)

    checks = sum(
        bool(v)
        for v in details.values()
    )

    maximum = CRITERIA_WEIGHTS[
        "contact_information"
    ]

    score = maximum * (
        checks / len(details)
    )

    evidence = [
        name
        for name, present in [
            (
                "Email detected",
                details["email"]
            ),
            (
                "Phone number detected",
                details["phone"]
            ),
            (
                "LinkedIn detected",
                details["linkedin"]
            ),
            (
                "GitHub detected",
                details["github"]
            ),
            (
                "Portfolio link detected",
                details["portfolio"]
            ),
        ]
        if present
    ]

    missing = [
        name
        for name, present in [
            (
                "email",
                details["email"]
            ),
            (
                "phone",
                details["phone"]
            ),
            (
                "LinkedIn",
                details["linkedin"]
            ),
            (
                "GitHub/portfolio",
                details["github"]
                or details["portfolio"]
            ),
        ]
        if not present
    ]

    explanation = (
        "Contact details are machine-readable and include "
        + (
            ", ".join(evidence)
            if evidence
            else "little usable contact information"
        )
        + "."
    )

    if missing:
        explanation += (
            " Consider adding: "
            + ", ".join(missing)
            + "."
        )

    return (
        make_result(
            "contact_information",
            score,
            evidence,
            explanation
        ),
        details
    )


# ============================================================
# 2. RESUME PARSING
# ============================================================

def check_resume_parsing(text, pages):

    wc = word_count(text)

    weird_chars = len(
        re.findall(r"[�]", text)
    )

    repeated_spaces = len(
        re.findall(r" {3,}", text)
    )

    lines = split_lines(text)

    quality = 1.0

    if wc < 120:
        quality -= 0.35

    elif wc < 200:
        quality -= 0.15

    if weird_chars:
        quality -= min(
            0.30,
            weird_chars / 20
        )

    if not lines:
        quality = 0

    score = (
        CRITERIA_WEIGHTS["resume_parsing"]
        * clamp(quality)
    )

    evidence = [
        f"Extracted approximately {wc} words",
        f"Detected {pages or 1} page(s)",
        f"Detected {len(lines)} non-empty text lines",
    ]

    if weird_chars:
        evidence.append(
            f"{weird_chars} replacement characters detected"
        )

    if repeated_spaces:
        evidence.append(
            f"{repeated_spaces} excessive spacing patterns detected"
        )

    explanation = (
        "The resume text appears sufficiently "
        "extractable for automated analysis."
        if quality >= 0.80
        else
        "The extracted text has quality limitations "
        "that may reduce scoring reliability."
    )

    return make_result(
        "resume_parsing",
        score,
        evidence,
        explanation
    )


# ============================================================
# 3. JOB DESCRIPTION MATCHING
# ============================================================

def check_job_description_matching(text):

    return na_result(
        "job_description_matching",
        "Project 1 is resume-only. "
        "Job Description matching is reserved for Project 2."
    )


# ============================================================
# 4. KEYWORD MATCHING
# ============================================================

def check_keyword_matching(text):

    skills = find_skills(text)

    maximum = CRITERIA_WEIGHTS[
        "keyword_matching"
    ]

    count = len(skills)

    score = maximum * clamp(
        count / 15
    )

    evidence = [
        f"{count} recognized skills/technical keywords detected"
    ]

    if skills:
        evidence.append(
            "Examples: "
            + ", ".join(skills[:10])
        )

    explanation = (
        "This is a resume-only keyword coverage check. "
        "It measures recognizable professional vocabulary, "
        "not job-description matching."
    )

    return make_result(
        "keyword_matching",
        score,
        evidence,
        explanation
    )


# ============================================================
# 5. SKILLS ANALYSIS
# ============================================================

def check_skills_analysis(text):

    skills = find_skills(text)

    lower = text.lower()

    section_present = any(
        alias in lower
        for alias in SECTION_ALIASES["skills"]
    )

    count = len(skills)

    base = clamp(
        count / 12
    )

    if section_present:
        base = min(
            1.0,
            base + 0.15
        )

    score = (
        CRITERIA_WEIGHTS["skills_analysis"]
        * base
    )

    evidence = [
        f"{count} recognized skills detected",

        (
            "Dedicated skills section detected"
            if section_present
            else
            "Dedicated skills section not clearly detected"
        ),
    ]

    if skills:
        evidence.append(
            "Skills: "
            + ", ".join(skills[:15])
        )

    explanation = (
        "Skills are clearly represented and include "
        "several recognizable technical competencies."
        if base >= 0.80
        else
        "The skills inventory could be made more complete "
        "and clearly structured."
    )

    return make_result(
        "skills_analysis",
        score,
        evidence,
        explanation
    )


# ============================================================
# 6. RESUME SECTIONS
# ============================================================

def check_resume_sections(text):

    sections, aliases = detect_sections(text)

    required = [
        "summary",
        "education",
        "skills",
        "projects"
    ]

    useful_optional = [
        "experience",
        "certifications",
        "achievements"
    ]

    required_found = sum(
        sections[s]
        for s in required
    )

    optional_found = sum(
        sections[s]
        for s in useful_optional
    )

    ratio = (
        required_found / len(required)
    ) * 0.75

    ratio += (
        optional_found / len(useful_optional)
    ) * 0.25

    score = (
        CRITERIA_WEIGHTS["resume_sections"]
        * ratio
    )

    evidence = [
        f"{sum(sections.values())} recognized resume sections detected"
    ]

    evidence.extend(
        f"{CRITERION_LABELS.get(s, s.title())}: "
        f"{'present' if present else 'not detected'}"
        for s, present in sections.items()
        if s in required + useful_optional
    )

    explanation = (
        "The resume has a recognizable, "
        "structured section layout."
        if ratio >= 0.80
        else
        "Several standard resume sections "
        "could be clearer or more complete."
    )

    return make_result(
        "resume_sections",
        score,
        evidence,
        explanation
    )


# ============================================================
# 7. EXPERIENCE MATCHING
# ============================================================

def check_experience_matching(text):

    sections, _ = detect_sections(text)

    lower = text.lower()

    experience_terms = [
        "intern",
        "internship",
        "developer",
        "engineer",
        "analyst",
        "assistant",
        "trainee",
        "research",
        "freelance",
        "volunteer"
    ]

    found_terms = [
        term
        for term in experience_terms
        if re.search(
            r"(?<!\w)"
            + re.escape(term)
            + r"(?!\w)",
            lower
        )
    ]

    years = extract_years(text)

    bullets = bullet_lines(text)

    if sections["experience"] or found_terms:

        evidence = [
            "Experience/internship terminology detected",
            f"{len(bullets)} bullet-style lines detected",
        ]

        if years:
            evidence.append(
                "Date/year information detected"
            )

        ratio = (
            0.85
            if bullets
            else 0.65
        )

        if years:
            ratio = min(
                1.0,
                ratio + 0.10
            )

        explanation = (
            "Experience-related evidence is present. "
            "For a fresher, internships, training, "
            "academic work, and relevant practical "
            "experience can provide useful evidence."
        )

    else:

        evidence = [
            "No clear professional experience/internship evidence detected"
        ]

        # Do not punish a fresher to zero merely
        # for lacking employment.
        ratio = 0.55

        explanation = (
            "No clear work experience was detected. "
            "This is not treated as a failure for a fresher; "
            "adding internships or practical experience "
            "can strengthen the evidence."
        )

    score = (
        CRITERIA_WEIGHTS["experience_matching"]
        * ratio
    )

    return make_result(
        "experience_matching",
        score,
        evidence,
        explanation
    )


# ============================================================
# 8. PROJECT RELEVANCE
# ============================================================

def check_project_relevance(text):

    sections, _ = detect_sections(text)

    lower = text.lower()

    project_terms = [
        "project",
        "developed",
        "built",
        "implemented",
        "using",
        "technology",
        "github",
        "dataset",
        "model",
        "application"
    ]

    term_count = sum(
        bool(
            re.search(
                r"(?<!\w)"
                + re.escape(term)
                + r"(?!\w)",
                lower
            )
        )
        for term in project_terms
    )

    bullets = bullet_lines(text)

    ratio = 0.0

    if sections["projects"]:
        ratio += 0.50

    if term_count >= 4:
        ratio += 0.25

    if bullets:
        ratio += 0.15

    if len(find_skills(text)) >= 4:
        ratio += 0.10

    ratio = clamp(ratio)

    evidence = [
        (
            "Projects section detected"
            if sections["projects"]
            else
            "Projects section not clearly detected"
        ),
        f"{term_count} project-context indicators detected",
        f"{len(bullets)} bullet-style lines detected",
    ]

    explanation = (
        "Projects contain useful evidence of applied work."
        if ratio >= 0.80
        else
        "Projects could show clearer responsibilities, "
        "technologies, and measurable outcomes."
    )

    return make_result(
        "project_relevance",
        CRITERIA_WEIGHTS["project_relevance"] * ratio,
        evidence,
        explanation
    )


# ============================================================
# 9. EDUCATION
# ============================================================

def check_education(text):

    lower = text.lower()

    found = [
        term.strip()
        for term in DEGREE_TERMS
        if re.search(
            re.escape(term.strip()),
            lower
        )
    ]

    years = extract_years(text)

    sections, _ = detect_sections(text)

    ratio = 0.0

    if sections["education"]:
        ratio += 0.55

    if found:
        ratio += 0.30

    if years:
        ratio += 0.15

    ratio = clamp(ratio)

    evidence = [
        (
            "Education section detected"
            if sections["education"]
            else
            "Education section not clearly detected"
        ),

        (
            "Degree/academic terminology detected"
            if found
            else
            "Degree terminology not clearly detected"
        ),

        (
            "Academic year/date detected"
            if years
            else
            "Academic date not clearly detected"
        ),
    ]

    explanation = (
        "Education information is reasonably machine-readable."
        if ratio >= 0.80
        else
        "Make the degree, specialization, institution, "
        "and graduation year easy to identify."
    )

    return make_result(
        "education",
        CRITERIA_WEIGHTS["education"] * ratio,
        evidence,
        explanation
    )


# ============================================================
# 10. CERTIFICATIONS
# ============================================================

def check_certifications(text):

    sections, _ = detect_sections(text)

    lower = text.lower()

    certification_terms = [
        "certified",
        "certificate",
        "certification",
        "coursera",
        "udemy",
        "nptel",
        "aws certified",
        "microsoft certified",
        "google certificate"
    ]

    found = [
        term
        for term in certification_terms
        if term in lower
    ]

    if sections["certifications"] or found:

        ratio = 1.0

        evidence = [
            "Certification evidence detected"
        ]

        if sections["certifications"]:
            evidence.append(
                "Certifications section detected"
            )

        explanation = (
            "Relevant certification information is present."
        )

    else:

        # Certifications are optional,
        # especially for students/freshers.
        ratio = 0.50

        evidence = [
            "No certification section/evidence detected"
        ]

        explanation = (
            "No certifications were detected. "
            "This is treated as optional, "
            "not as an automatic failure."
        )

    return make_result(
        "certifications",
        CRITERIA_WEIGHTS["certifications"] * ratio,
        evidence,
        explanation
    )


# ============================================================
# 11. ACHIEVEMENTS & METRICS
# ============================================================

def check_achievements_metrics(text):

    metrics = extract_metrics(text)

    sections, _ = detect_sections(text)

    ratio = clamp(
        len(metrics) / 5
    )

    if sections["achievements"]:
        ratio = min(
            1.0,
            ratio + 0.20
        )

    evidence = [
        f"{len(metrics)} measurable result patterns detected"
    ]

    if metrics:
        evidence.append(
            "Examples: "
            + ", ".join(metrics[:8])
        )

    explanation = (
        "The resume includes measurable evidence."
        if ratio >= 0.80
        else
        "Add measurable outcomes such as percentages, "
        "scale, time saved, accuracy, users, or records."
    )

    return make_result(
        "achievements_metrics",
        CRITERIA_WEIGHTS["achievements_metrics"] * ratio,
        evidence,
        explanation
    )


# ============================================================
# 12. ACTION VERBS
# ============================================================

def check_action_verbs(text):

    actions = count_action_verbs(text)

    ratio = clamp(
        len(actions) / 10
    )

    evidence = [
        f"{len(actions)} distinct action verbs detected"
    ]

    if actions:
        evidence.append(
            "Examples: "
            + ", ".join(actions[:12])
        )

    explanation = (
        "Resume bullets use a useful range "
        "of action-oriented language."
        if ratio >= 0.80
        else
        "Use precise action verbs such as developed, "
        "implemented, analyzed, optimized, and deployed."
    )

    return make_result(
        "action_verbs",
        CRITERIA_WEIGHTS["action_verbs"] * ratio,
        evidence,
        explanation
    )


# ============================================================
# 13. ATS FORMATTING
# ============================================================

def check_ats_formatting(text):

    sections, _ = detect_sections(text)

    lines = split_lines(text)

    long_lines = [
        line
        for line in lines
        if len(line) > 180
    ]

    bullets = bullet_lines(text)

    urls = URL_PATTERN.findall(text)

    standard_heading_count = sum(
        sections[s]
        for s in [
            "summary",
            "experience",
            "education",
            "skills",
            "projects"
        ]
    )

    ratio = 0.0

    ratio += (
        0.45
        * clamp(
            standard_heading_count / 4
        )
    )

    ratio += (
        0.20
        if bullets
        else 0
    )

    ratio += (
        0.15
        if not long_lines
        else 0.0
    )

    ratio += (
        0.10
        if len(urls) <= 8
        else 0.0
    )

    ratio += (
        0.10
        if len(text) >= 500
        else 0.0
    )

    evidence = [
        f"{standard_heading_count}/5 standard headings detected",
        f"{len(bullets)} bullet-style lines detected",
        f"{len(long_lines)} very long lines detected",
        f"{len(urls)} URLs detected",
    ]

    explanation = (
        "The extracted text has several ATS-friendly "
        "structural signals."
        if ratio >= 0.80
        else
        "Use standard headings, simple text structure, "
        "consistent bullets, and keep important "
        "information out of images."
    )

    return make_result(
        "ats_formatting",
        CRITERIA_WEIGHTS["ats_formatting"]
        * clamp(ratio),
        evidence,
        explanation
    )


# ============================================================
# 14. READABILITY
# ============================================================

def check_readability(text):

    wc = word_count(text)

    sentences = sentence_count(text)

    avg_sentence_words = (
        wc / sentences
        if sentences
        else wc
    )

    lines = split_lines(text)

    overly_long = sum(
        1
        for line in lines
        if len(line.split()) > 45
    )

    ratio = 1.0

    if wc < 200:
        ratio -= 0.20

    if avg_sentence_words > 35:
        ratio -= 0.20

    if overly_long:
        ratio -= min(
            0.25,
            overly_long * 0.05
        )

    ratio = clamp(ratio)

    evidence = [
        f"Approximately {wc} words",
        f"Average sentence length approximately "
        f"{avg_sentence_words:.1f} words",
        f"{overly_long} very long text lines detected",
    ]

    explanation = (
        "The text is reasonably concise and readable "
        "for automated and human review."
        if ratio >= 0.80
        else
        "Shorten long bullets and use concise, "
        "result-focused statements."
    )

    return make_result(
        "readability",
        CRITERIA_WEIGHTS["readability"] * ratio,
        evidence,
        explanation
    )


# ============================================================
# 15. KEYWORD CONTEXT
# ============================================================

def check_keyword_context(text):

    skills = find_skills(text)

    lower = text.lower()

    contextual_hits = 0
    weak_hits = 0

    examples = []

    for skill in skills:

        pattern = re.compile(
            r".{0,100}"
            + re.escape(skill.lower())
            + r".{0,100}",
            re.I | re.S
        )

        matches = pattern.findall(lower)

        if not matches:
            continue

        contextual = False

        for match in matches[:3]:

            if any(
                verb in match
                for verb in [
                    "used",
                    "built",
                    "developed",
                    "implemented",
                    "designed",
                    "analyzed",
                    "created",
                    "deployed",
                    "worked",
                    "trained",
                    "applied"
                ]
            ):
                contextual = True
                break

        if contextual:

            contextual_hits += 1

            if len(examples) < 5:
                examples.append(skill)

        else:
            weak_hits += 1

    total = len(skills)

    if total:
        ratio = (
            contextual_hits
            + 0.35 * weak_hits
        ) / total

    else:
        ratio = 0

    evidence = [
        f"{contextual_hits} skills found in meaningful action/context",
        f"{weak_hits} skills found with weak context",
    ]

    if examples:
        evidence.append(
            "Context-rich examples: "
            + ", ".join(examples)
        )

    explanation = (
        "Many detected skills are supported "
        "by surrounding action/context."
        if ratio >= 0.75
        else
        "Skills are stronger when supported by projects, "
        "experience, or specific actions rather than "
        "appearing only as a list."
    )

    return make_result(
        "keyword_context",
        CRITERIA_WEIGHTS["keyword_context"]
        * clamp(ratio),
        evidence,
        explanation
    )


# ============================================================
# 16. MISSING REQUIREMENTS
# ============================================================

def check_missing_requirements(text):

    return na_result(
        "missing_requirements",
        "A missing-requirements comparison requires "
        "a Job Description, so it is reserved for Project 2."
    )


# ============================================================
# 17. POTENTIAL RED FLAGS
# ============================================================

def check_red_flags(text):

    lower = text.lower()

    flags = []

    if re.search(
        r"\b(lorem ipsum|test resume|sample resume)\b",
        lower
    ):
        flags.append(
            "Placeholder/sample text detected"
        )

    if EMAIL_PATTERN.search(text) is None:
        flags.append(
            "No email detected"
        )

    # Excessive keyword repetition can be a warning,
    # not a reason to reject by itself.
    skills = find_skills(text)

    for skill in skills:

        count = len(
            re.findall(
                r"(?<!\w)"
                + re.escape(skill.lower())
                + r"(?!\w)",
                lower
            )
        )

        if count >= 15:
            flags.append(
                f"Possible keyword overuse: {skill}"
            )

    if len(
        re.findall(
            r"[^\x00-\x7F]",
            text
        )
    ) > max(
        50,
        len(text) * 0.10
    ):
        flags.append(
            "Unusual character density may affect parsing"
        )

    years = [
        int(y)
        for y in re.findall(
            r"\b(?:19|20)\d{2}\b",
            text
        )
    ]

    if (
        len(years) >= 3
        and max(years) - min(years) > 80
    ):
        flags.append(
            "Unusual date range detected"
        )

    maximum = CRITERIA_WEIGHTS[
        "potential_red_flags"
    ]

    ratio = clamp(
        1 - len(flags) / 5
    )

    evidence = (
        ["No major automated red flags detected"]
        if not flags
        else flags
    )

    explanation = (
        "No major automated red flags were detected."
        if not flags
        else
        "Potential issues were detected. "
        "Review them manually before changing the resume."
    )

    return make_result(
        "potential_red_flags",
        maximum * ratio,
        evidence,
        explanation
    )


# ============================================================
# 18. CONTENT CONSISTENCY
# ============================================================

def check_content_consistency(text):

    sections, _ = detect_sections(text)

    years = [
        int(y)
        for y in re.findall(
            r"\b(?:19|20)\d{2}\b",
            text
        )
    ]

    duplicate_lines = [
        line.lower()
        for line in split_lines(text)
        if len(line.split()) >= 5
    ]

    counts = Counter(
        duplicate_lines
    )

    duplicates = [
        line
        for line, count in counts.items()
        if count > 1
    ]

    ratio = 1.0

    if len(duplicates) > 2:
        ratio -= 0.25

    if years and (
        min(years) < 1950
        or max(years) > 2100
    ):
        ratio -= 0.20

    ratio = clamp(ratio)

    evidence = [
        f"{len(duplicates)} repeated substantial lines detected",
        f"{len(years)} year/date tokens detected",
        f"{sum(sections.values())} section types detected",
    ]

    explanation = (
        "No major internal consistency problems "
        "were detected by the automated checks."
        if ratio >= 0.80
        else
        "Review repeated content and dates for consistency."
    )

    return make_result(
        "content_consistency",
        CRITERIA_WEIGHTS["content_consistency"]
        * ratio,
        evidence,
        explanation
    )


# ============================================================
# SCORE ENGINE
# ============================================================

def calculate_ats_score(text, pages=0):
    """
    Runs all 18 checks.

    JD-dependent criteria are transparently N/A
    for Project 1.

    Their weights are redistributed across the
    applicable criteria proportionally.
    """

    results = []

    # 1
    contact, contact_details = (
        check_contact_information(text)
    )

    results.append(contact)

    # 2
    results.append(
        check_resume_parsing(
            text,
            pages
        )
    )

    # 3
    results.append(
        check_job_description_matching(text)
    )

    # 4
    results.append(
        check_keyword_matching(text)
    )

    # 5
    results.append(
        check_skills_analysis(text)
    )

    # 6
    results.append(
        check_resume_sections(text)
    )

    # 7
    results.append(
        check_experience_matching(text)
    )

    # 8
    results.append(
        check_project_relevance(text)
    )

    # 9
    results.append(
        check_education(text)
    )

    # 10
    results.append(
        check_certifications(text)
    )

    # 11
    results.append(
        check_achievements_metrics(text)
    )

    # 12
    results.append(
        check_action_verbs(text)
    )

    # 13
    results.append(
        check_ats_formatting(text)
    )

    # 14
    results.append(
        check_readability(text)
    )

    # 15
    results.append(
        check_keyword_context(text)
    )

    # 16
    results.append(
        check_missing_requirements(text)
    )

    # 17
    results.append(
        check_red_flags(text)
    )

    # 18
    results.append(
        check_content_consistency(text)
    )

    # --------------------------------------------------------
    # Find applicable criteria
    # --------------------------------------------------------

    applicable = [
        result
        for result in results
        if result["status"] != "N/A"
    ]

    applicable_weight = sum(
        result["max_score"]
        for result in applicable
    )

    na_weight = (
        100 - applicable_weight
    )

    # --------------------------------------------------------
    # Redistribute N/A weight
    # --------------------------------------------------------

    final_score = 0.0

    for result in applicable:

        original_weight = result["max_score"]

        adjusted_weight = (
            original_weight
            / applicable_weight
            * 100
            if applicable_weight
            else 0
        )

        result["adjusted_weight"] = round(
            adjusted_weight,
            2
        )

        result["weighted_contribution"] = round(
            (
                result["score"]
                / original_weight
            )
            * adjusted_weight
            if original_weight
            else 0,
            2
        )

        final_score += (
            result["weighted_contribution"]
        )

    # N/A criteria contribute zero.
    for result in results:

        if result["status"] == "N/A":

            result["adjusted_weight"] = 0

            result["weighted_contribution"] = 0

    final_score = round(
        clamp(
            final_score / 100
        ) * 100,
        2
    )

    # --------------------------------------------------------
    # Confidence
    # --------------------------------------------------------

    parsing = next(
        r
        for r in results
        if r["key"] == "resume_parsing"
    )

    parsing_ratio = (
        parsing["score"]
        / parsing["max_score"]
        if parsing["max_score"]
        else 0
    )

    if parsing_ratio >= 0.80:
        confidence = "High"

    elif parsing_ratio >= 0.55:
        confidence = "Medium"

    else:
        confidence = "Low"

    return (
        final_score,
        results,
        confidence,
        contact_details
    )


# ============================================================
# ADVICE ENGINE
# ============================================================

def generate_advice(results, text):

    advice = []

    priority_order = {
        "High": 1,
        "Medium": 2,
        "Low": 3
    }

    for result in results:

        if result["status"] == "N/A":
            continue

        ratio = (
            result["score"]
            / result["max_score"]
            if result["max_score"]
            else 1
        )

        if ratio < 0.50:
            priority = "High"

        elif ratio < 0.75:
            priority = "Medium"

        else:
            priority = "Low"

        if (
            result["key"]
            == "contact_information"
            and ratio < 0.80
        ):
            advice.append({
                "priority": priority,
                "title": "Complete contact information",
                "advice":
                    "Add missing email, phone, "
                    "LinkedIn, GitHub, or portfolio details."
            })

        elif (
            result["key"]
            == "resume_sections"
            and ratio < 0.80
        ):
            advice.append({
                "priority": priority,
                "title": "Strengthen resume sections",
                "advice":
                    "Use clear headings such as Summary, "
                    "Skills, Projects, Education, and Experience."
            })

        elif (
            result["key"]
            == "skills_analysis"
            and ratio < 0.75
        ):
            advice.append({
                "priority": priority,
                "title": "Improve the skills section",
                "advice":
                    "List relevant skills you genuinely know "
                    "and support important skills elsewhere "
                    "in the resume."
            })

        elif (
            result["key"]
            == "project_relevance"
            and ratio < 0.75
        ):
            advice.append({
                "priority": priority,
                "title": "Strengthen projects",
                "advice":
                    "For each project, show your role, "
                    "technologies used, what you built, "
                    "and the result."
            })

        elif (
            result["key"]
            == "achievements_metrics"
            and ratio < 0.75
        ):
            advice.append({
                "priority": priority,
                "title": "Add measurable results",
                "advice":
                    "Where truthful, quantify impact using "
                    "percentages, users, records, accuracy, "
                    "time saved, or scale."
            })

        elif (
            result["key"]
            == "action_verbs"
            and ratio < 0.75
        ):
            advice.append({
                "priority": priority,
                "title": "Use stronger action verbs",
                "advice":
                    "Start bullets with precise verbs such as "
                    "developed, implemented, analyzed, "
                    "optimized, or deployed."
            })

        elif (
            result["key"]
            == "ats_formatting"
            and ratio < 0.75
        ):
            advice.append({
                "priority": priority,
                "title": "Improve ATS-friendly formatting",
                "advice":
                    "Use standard headings, simple text structure, "
                    "consistent bullets, and keep important "
                    "information out of images."
            })

        elif (
            result["key"]
            == "readability"
            and ratio < 0.75
        ):
            advice.append({
                "priority": priority,
                "title": "Improve readability",
                "advice":
                    "Shorten long bullets and remove repeated "
                    "or low-value wording."
            })

        elif (
            result["key"]
            == "keyword_context"
            and ratio < 0.75
        ):
            advice.append({
                "priority": priority,
                "title": "Support skills with context",
                "advice":
                    "Show important skills in projects or "
                    "experience with a specific action rather "
                    "than only listing them."
            })

        elif (
            result["key"]
            == "potential_red_flags"
            and ratio < 0.75
        ):
            advice.append({
                "priority": priority,
                "title": "Review potential red flags",
                "advice":
                    "Check the flagged items manually before "
                    "making changes; automated flags are "
                    "signals, not proof of an error."
            })

    if not advice:
        advice.append({
            "priority": "Low",
            "title": "Maintain the current structure",
            "advice":
                "The automated checks found no major "
                "improvement area. Continue tailoring "
                "the resume truthfully for each target "
                "role in Project 2."
        })

    advice.sort(
        key=lambda x:
        priority_order.get(
            x["priority"],
            9
        )
    )

    return advice


def get_skills_to_improve(found_skills):

    recommended = [
        "python",
        "sql",
        "git",
        "data analysis",
        "machine learning",
        "pandas",
        "numpy",
        "data visualization"
    ]

    found_lower = {
        skill.lower()
        for skill in found_skills
    }

    return [
        skill
        for skill in recommended
        if skill.lower()
        not in found_lower
    ][:6]


# ============================================================
# ROUTES
# ============================================================

@app.route("/")
def home():
    return render_template(
        "index.html"
    )


@app.route(
    "/upload",
    methods=["POST"]
)
def upload_resume():

    if "resume" not in request.files:

        return jsonify({
            "success": False,
            "message":
                "No resume file received."
        }), 400

    file = request.files["resume"]

    if file.filename == "":

        return jsonify({
            "success": False,
            "message":
                "No file selected."
        }), 400

    if not allowed_file(
        file.filename
    ):

        return jsonify({
            "success": False,
            "message":
                "Only PDF, DOC and DOCX files are allowed."
        }), 400

    filename = secure_filename(
        file.filename
    )

    if not filename:

        return jsonify({
            "success": False,
            "message":
                "Invalid filename."
        }), 400

    filepath = os.path.join(
        app.config["UPLOAD_FOLDER"],
        filename
    )

    file.save(filepath)

    try:

        resume_text, pages = (
            extract_resume_text(
                filepath,
                filename
            )
        )

        resume_text = normalize_text(
            resume_text
        )

        if not resume_text:

            return jsonify({
                "success": False,
                "message":
                    "Could not read the resume. "
                    "Please use a text-based PDF "
                    "or DOCX file. Legacy DOC files "
                    "must be converted first."
            }), 400

        (
            ats_score,
            results,
            confidence,
            contact_details
        ) = calculate_ats_score(
            resume_text,
            pages
        )

        skills = find_skills(
            resume_text
        )

        sections, _ = detect_sections(
            resume_text
        )

        actions = count_action_verbs(
            resume_text
        )

        metrics = extract_metrics(
            resume_text
        )

        advice = generate_advice(
            results,
            resume_text
        )

        return jsonify({

            "success": True,

            "project":
                "Project 1 - Resume-only ATS Checker",

            "filename":
                filename,

            # ------------------------------------------------
            # Main score
            # ------------------------------------------------

            "ats_score":
                ats_score,

            "score_out_of":
                100,

            "score_confidence":
                confidence,

            # ------------------------------------------------
            # All 18 checks
            # ------------------------------------------------

            "criteria_count":
                len(results),

            "criteria_completed":
                len(results),

            "criteria":
                results,

            # ------------------------------------------------
            # Compatibility with frontend
            # ------------------------------------------------

            "score_breakdown": {
                r["key"]: r["score"]
                for r in results
            },

            "top_matching_skills":
                skills,

            "skills_to_improve":
                get_skills_to_improve(
                    skills
                ),

            "improvement_advice":
                advice,

            # ------------------------------------------------
            # Analysis data
            # ------------------------------------------------

            "analysis": {

                "sections":
                    sections,

                "contact_information":
                    contact_details,

                "action_words":
                    actions,

                "quantifiable_metrics":
                    metrics,

                "word_count":
                    word_count(
                        resume_text
                    ),

                "page_count":
                    pages,
            },

            # ------------------------------------------------
            # Project notes
            # ------------------------------------------------

            "notes": [

                "Project 1 is resume-only.",

                "Job Description Matching and "
                "Missing Requirements are N/A here.",

                "No score is fabricated for "
                "unavailable Job Description evidence.",

                "The final score is normalized "
                "to 100 across applicable criteria."
            ]
        })

    except Exception as exc:

        app.logger.exception(
            "Resume analysis failed"
        )

        return jsonify({
            "success": False,
            "message":
                "Resume analysis failed safely. "
                "Please try another file."
        }), 500

    finally:

        # Remove uploaded resume after analysis.
        try:

            if os.path.exists(filepath):
                os.remove(filepath)

        except OSError:
            pass


# ============================================================
# ERROR HANDLERS
# ============================================================

@app.errorhandler(413)
def file_too_large(error):

    return jsonify({
        "success": False,
        "message":
            "File is too large. Maximum size is 10 MB."
    }), 413


# ============================================================
# RUN APPLICATION
# ============================================================

if __name__ == "__main__":
    app.run(
        debug=True
    )